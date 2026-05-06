"""
parsers/assessor_guide_parser.py

Starter parser for 'Assessor Guide' style DOCX files (project/rubric/task layouts).

This module provides a pragmatic, heuristic-based extractor that:
- groups sequential paragraphs and tables into high-level sections (Project/Stage/Task/Rubric/Appendix/Instructions)
- extracts tables that look like rubrics into lists of criteria (header -> value mapping)
- preserves red-formatted text as assessor feedback entries

The goal is to provide a structured representation we can use to author assessment items
and rubrics into the rest of the system. It intentionally favors readability and small scope.
"""

from __future__ import annotations

import re
from typing import List, Dict, Any

from docx import Document

from parsers.docx_extractor import (
    iter_block_items,
    paragraph_text_and_is_red,
    textbox_texts_in_paragraph,
)


# Heuristics to detect section starts
PROJECT_RE = re.compile(r"^Project\b|^Project\s*\d", re.I)
STAGE_RE = re.compile(r"^Stage\b|^Stage\s*\d", re.I)
TASK_RE = re.compile(r"^Task\b|^Task\s*\d", re.I)
RUBRIC_RE = re.compile(r"\bRubric\b", re.I)
APPENDIX_RE = re.compile(r"^Appendix\b", re.I)
ASSESSOR_INSTR_RE = re.compile(r"assessor instructions|instructions for assessors|assessor is to", re.I)


def _text_from_cell(cell) -> str:
    parts = []
    for p in cell.paragraphs:
        t, _ = paragraph_text_and_is_red(p)
        if t:
            parts.append(t)
    return "\n".join(parts).strip()


def _parse_table_as_dicts(table) -> Dict[str, Any]:
    """Return parsed table as {'headers': [...], 'rows': [dict,...]} using first row as header when possible."""
    rows = list(table.rows)
    if not rows:
        return {"headers": [], "rows": []}

    header_cells = rows[0].cells
    headers = [_text_from_cell(c) or f"col{i}" for i, c in enumerate(header_cells)]

    out_rows: List[Dict[str, str]] = []
    for row in rows[1:]:
        cells = row.cells
        rowd: Dict[str, str] = {}
        for i, c in enumerate(cells):
            h = headers[i] if i < len(headers) else f"col{i}"
            rowd[h] = _text_from_cell(c)
        out_rows.append(rowd)

    return {"headers": headers, "rows": out_rows}


def parse_assessor_guide(docx_path: str) -> Dict[str, Any]:
    """Parse the assessor guide into sections, rubrics and feedback.

    Returns a dict with keys:
    - sections: list of {type, title, paragraphs: [str], tables: [parsed_table_dicts]}
    - rubrics: list of parsed rubric tables
    - feedback_lines: list of red-formatted strings (assessor feedback)
    """
    doc = Document(docx_path)
    sections: List[Dict[str, Any]] = []
    rubrics: List[Dict[str, Any]] = []
    feedback_lines: List[str] = []

    current = {"type": "body", "title": None, "paragraphs": [], "tables": []}

    def push_current():
        nonlocal current
        if current and (current.get("paragraphs") or current.get("tables")):
            sections.append(current)
        current = {"type": "body", "title": None, "paragraphs": [], "tables": []}

    for block in iter_block_items(doc):
        if hasattr(block, "rows"):
            # table
            parsed = _parse_table_as_dicts(block)
            current["tables"].append(parsed)

            # collect red feedback inside table cells
            for row in block.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        t, is_red = paragraph_text_and_is_red(p)
                        if t and is_red:
                            feedback_lines.append(t)

            # heuristics: if table header contains rubric-like words, capture as rubric
            header_flat = " ".join(parsed.get("headers", []))
            if re.search(r"criteria|element|satisfactory|not yet|rubric|marking criteria", header_flat, re.I):
                rubrics.append({"title": current.get("title"), "table": parsed})
            continue

        # paragraph
        t, is_red = paragraph_text_and_is_red(block)
        if not t:
            # still check for textbox content inside paragraph
            for t2, red2 in textbox_texts_in_paragraph(block):
                if t2:
                    current["paragraphs"].append(t2)
                    if red2:
                        feedback_lines.append(t2)
            continue

        # capture red lines separately
        if is_red:
            feedback_lines.append(t)

        # detect section headers
        if PROJECT_RE.search(t):
            push_current()
            current = {"type": "project", "title": t, "paragraphs": [], "tables": []}
            continue
        if STAGE_RE.search(t):
            push_current()
            current = {"type": "stage", "title": t, "paragraphs": [], "tables": []}
            continue
        if TASK_RE.search(t):
            push_current()
            current = {"type": "task", "title": t, "paragraphs": [], "tables": []}
            continue
        if RUBRIC_RE.search(t):
            push_current()
            current = {"type": "rubric", "title": t, "paragraphs": [], "tables": []}
            continue
        if APPENDIX_RE.search(t):
            push_current()
            current = {"type": "appendix", "title": t, "paragraphs": [], "tables": []}
            continue
        if ASSESSOR_INSTR_RE.search(t):
            # these often are red; but if not, mark as assessor section
            push_current()
            current = {"type": "assessor_instructions", "title": t, "paragraphs": [], "tables": []}
            continue

        # otherwise append to current
        current["paragraphs"].append(t)

    # final flush
    push_current()

    return {
        "sections": sections,
        "rubrics": rubrics,
        "feedback_lines": feedback_lines,
    }


if __name__ == "__main__":
    # simple CLI for quick debugging
    import json
    import sys

    path = sys.argv[1] if len(sys.argv) > 1 else "rubics/BSBHRM613 Project - Assessor Guide.docx"
    out = parse_assessor_guide(path)
    print(json.dumps({"sections": len(out["sections"]), "rubrics": len(out["rubrics"]), "feedback_lines": len(out["feedback_lines"])}, indent=2))
